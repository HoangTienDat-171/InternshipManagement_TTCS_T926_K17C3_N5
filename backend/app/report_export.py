"""Excel and PDF presentation of already-filtered reporting rows."""

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


@dataclass(frozen=True)
class ReportColumn:
    label: str
    key: str
    excel_width: int
    pdf_width: int


def _safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def make_excel(title, filters, summary, columns, rows):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Báo cáo")
    for index, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = column.excel_width
    sheet.append([title])
    sheet.append(["Xuất lúc", datetime.now().astimezone().strftime("%d/%m/%Y %H:%M %Z")])
    for label, value in filters:
        sheet.append([label, _safe_cell(value)])
    for label, value in summary:
        sheet.append([label, _safe_cell(value)])
    sheet.append([])
    header = []
    for column in columns:
        cell = WriteOnlyCell(sheet, value=column.label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        header.append(cell)
    sheet.append(header)
    for row in rows:
        sheet.append([_safe_cell(row.get(column.key)) for column in columns])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _font_name():
    name = "IMSReportDejaVu"
    if name in pdfmetrics.getRegisteredFontNames():
        return name
    candidates = (
        Path("C:/Windows/Fonts/DejaVuSans.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
    )
    for candidate in candidates:
        if candidate.is_file():
            pdfmetrics.registerFont(TTFont(name, str(candidate)))
            return name
    raise RuntimeError("Máy chủ thiếu font Unicode để xuất PDF tiếng Việt.")


def make_pdf(title, filters, summary, columns, rows):
    font = _font_name()
    output = BytesIO()
    document = canvas.Canvas(output, pagesize=landscape(A4), pageCompression=1)
    document.setTitle(title)
    width, height = landscape(A4)
    left = 34
    right = width - 34
    bottom = 42
    table_width = right - left
    total_column_width = sum(column.pdf_width for column in columns)
    widths = [table_width * column.pdf_width / total_column_width for column in columns]
    x_positions = [left]
    for col_width in widths:
        x_positions.append(x_positions[-1] + col_width)
    page = 0
    y = 0

    def new_page():
        nonlocal page, y
        if page:
            document.showPage()
        page += 1
        document.setFillColor(colors.HexColor("#4F46E5"))
        document.setFont(font, 15)
        document.drawString(left, height - 39, title)
        document.setFillColor(colors.HexColor("#475569"))
        document.setFont(font, 8)
        document.drawRightString(right, height - 38, f"Trang {page}")
        y = height - 60
        if page == 1:
            metadata = [("Xuất lúc", datetime.now().astimezone().strftime("%d/%m/%Y %H:%M")), *filters, *summary]
            for label, value in metadata:
                lines = simpleSplit(f"{label}: {value}", font, 8, table_width)
                for line in lines:
                    document.drawString(left, y, line)
                    y -= 12
            y -= 8
        document.setFillColor(colors.HexColor("#3730A3"))
        document.rect(left, y - 23, table_width, 25, fill=1, stroke=0)
        document.setFillColor(colors.white)
        document.setFont(font, 8)
        for index, column in enumerate(columns):
            document.drawString(x_positions[index] + 4, y - 13, column.label)
        y -= 25

    new_page()
    for row_number, row in enumerate(rows):
        wrapped = [simpleSplit(str(row.get(column.key) or ""), font, 8, widths[index] - 8) or [""]
                   for index, column in enumerate(columns)]
        line_number = 0
        max_lines = max(map(len, wrapped))
        while line_number < max_lines:
            available_lines = int((y - bottom - 8) // 11)
            if available_lines < 1:
                new_page()
                available_lines = int((y - bottom - 8) // 11)
            lines_to_draw = min(max_lines - line_number, available_lines)
            row_height = max(20, lines_to_draw * 11 + 8)
            if row_number % 2 == 0:
                document.setFillColor(colors.HexColor("#F5F6FF"))
                document.rect(left, y - row_height, table_width, row_height, fill=1, stroke=0)
            document.setFillColor(colors.HexColor("#1E293B"))
            document.setFont(font, 8)
            for index, lines in enumerate(wrapped):
                for offset, line in enumerate(lines[line_number:line_number + lines_to_draw]):
                    document.drawString(x_positions[index] + 4, y - 12 - offset * 11, line)
            y -= row_height
            line_number += lines_to_draw
    document.save()
    return output.getvalue()
