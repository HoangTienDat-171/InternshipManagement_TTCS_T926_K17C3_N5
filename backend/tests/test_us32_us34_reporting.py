"""Reporting checks against an isolated, in-memory SQLite fixture."""

import sqlite3
import unittest
from datetime import date
from io import BytesIO
from types import SimpleNamespace

from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader

from app.reporting_service import InternStatisticsFilters, InternStatisticsService
from app.routes import reporting_routes


def request_for(role):
    return SimpleNamespace(state=SimpleNamespace(current_user={"vai_tro": role}))


class ReportingTests(unittest.TestCase):
    def setUp(self):
        self.db = sqlite3.connect(":memory:")
        self.db.row_factory = sqlite3.Row
        self.db.executescript("""
            CREATE TABLE PHONG_BAN (ma_phong_ban INTEGER PRIMARY KEY, ten_phong_ban TEXT);
            CREATE TABLE TRUONG_DAI_HOC (ma_truong INTEGER PRIMARY KEY, ten_truong TEXT);
            CREATE TABLE NGUOI_DUNG (
                ma_nguoi_dung INTEGER PRIMARY KEY, ma_phong_ban INTEGER,
                ho_ten TEXT, email TEXT, vai_tro TEXT
            );
            CREATE TABLE HO_SO_THUC_TAP (
                ma_ho_so INTEGER PRIMARY KEY, ma_nguoi_dung INTEGER,
                ma_truong INTEGER, chuyen_nganh TEXT,
                trang_thai_xet_duyet TEXT, ngay_tao TEXT
            );
            CREATE TABLE CHUONG_TRINH_THUC_TAP (
                ma_chuong_trinh INTEGER PRIMARY KEY, ten_ct TEXT
            );
            CREATE TABLE UNG_TUYEN_CHUONG_TRINH (
                ma_ho_so INTEGER, ma_chuong_trinh INTEGER, trang_thai TEXT
            );
            INSERT INTO PHONG_BAN VALUES (1, 'Công nghệ'), (2, 'Thiết kế');
            INSERT INTO TRUONG_DAI_HOC VALUES (1, 'Đại học A'), (2, 'Đại học B');
            INSERT INTO NGUOI_DUNG VALUES
                (1, 1, 'Nguyễn An', 'an@example.test', 'ThucTapSinh'),
                (2, 1, 'Bình', 'binh@example.test', 'ThucTapSinh'),
                (3, 2, '=1+1', 'c@example.test', 'ThucTapSinh'),
                (4, 1, 'Mentor', 'm@example.test', 'Mentor');
            INSERT INTO HO_SO_THUC_TAP VALUES
                (11, 1, 1, 'Kỹ thuật phần mềm', 'DaDuyet', '2026-10-01 09:00:00'),
                (12, 2, 1, 'Thiết kế', 'DaDuyet', '2026-10-02 09:00:00'),
                (13, 3, NULL, 'Chưa cập nhật', 'ChoDuyet', '2026-10-03 09:00:00');
            INSERT INTO CHUONG_TRINH_THUC_TAP VALUES (10, 'Dự án A'), (20, 'Dự án B');
            INSERT INTO UNG_TUYEN_CHUONG_TRINH VALUES
                (11, 10, 'DaDuyet'), (11, 20, 'DaDuyet'),
                (12, 20, 'DaDuyet'), (13, 10, 'ChoDuyet');
        """)
        self.service = InternStatisticsService(self.db)

    def tearDown(self):
        self.db.close()

    def test_unique_profiles_and_participations_do_not_conflate(self):
        report = self.service.overview(InternStatisticsFilters(), 1, 10)
        self.assertEqual(report["summary"]["unique_interns"], 3)
        self.assertEqual(report["summary"]["approved_participations"], 3)
        self.assertEqual(sum(row["interns"] for row in report["schools"]), 3)
        self.assertEqual(sum(row["interns"] for row in report["majors"]), 3)
        self.assertEqual(report["summary"]["missing_school"], 1)
        self.assertEqual(report["summary"]["missing_major"], 1)

    def test_program_date_department_and_drilldown_share_cohort(self):
        filtered = self.service.overview(InternStatisticsFilters(
            program_id=20, department_id=1,
            date_from=date(2026, 10, 2), date_to=date(2026, 10, 2),
        ), 1, 10)
        self.assertEqual(filtered["summary"]["unique_interns"], 1)
        self.assertEqual(filtered["items"][0]["ma_ho_so"], 12)
        self.assertEqual(filtered["summary"]["approved_participations"], 1)
        school = self.service.overview(InternStatisticsFilters(school_id=1), 1, 10)
        self.assertEqual(school["schools"][0]["interns"], school["total_items"])
        missing = self.service.overview(InternStatisticsFilters(school_id=0), 1, 10)
        self.assertEqual(missing["items"][0]["ma_ho_so"], 13)
        self.assertEqual(missing["majors"][0]["interns"], missing["total_items"])

    def test_pagination_empty_and_invalid_date(self):
        first = self.service.overview(InternStatisticsFilters(), 1, 1)
        self.assertEqual(first["total_pages"], 3)
        self.assertEqual(len(first["items"]), 1)
        empty = self.service.overview(InternStatisticsFilters(program_id=99), 1, 10)
        self.assertEqual(empty["total_items"], 0)
        self.assertEqual(empty["total_pages"], 0)
        with self.assertRaises(HTTPException) as error:
            reporting_routes._filters(None, None, date(2026, 10, 3), date(2026, 10, 2), None, None, None)
        self.assertEqual(error.exception.status_code, 422)

    def test_role_guard_and_excel_export_match_dashboard(self):
        with self.assertRaises(HTTPException) as error:
            reporting_routes.intern_statistics_options(request_for("Mentor"), self.db)
        self.assertEqual(error.exception.status_code, 403)
        with self.assertRaises(HTTPException) as error:
            reporting_routes.export_intern_statistics(request_for("ThucTapSinh"), "xlsx", db=self.db)
        self.assertEqual(error.exception.status_code, 403)

        response = reporting_routes.export_intern_statistics(
            request_for("HR"), "xlsx", program_id=None, department_id=None,
            date_from=None, date_to=None, school_id=None, major=None, search=None, db=self.db,
        )
        self.assertEqual(response.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        workbook = load_workbook(BytesIO(response.body), read_only=True)
        values = list(workbook.active.values)
        data = [row for row in values if row and isinstance(row[0], int)]
        self.assertEqual(len(data), self.service.overview(InternStatisticsFilters(), 1, 10)["total_items"])
        self.assertTrue(any(row[1] == "'=1+1" for row in data))
        workbook.close()

        pdf_response = reporting_routes.export_intern_statistics(
            request_for("Admin"), "pdf", program_id=None, department_id=None,
            date_from=None, date_to=None, school_id=None, major=None, search=None, db=self.db,
        )
        reader = PdfReader(BytesIO(pdf_response.body))
        self.assertGreaterEqual(len(reader.pages), 1)
        text = "\n".join(page.extract_text() for page in reader.pages)
        self.assertIn("Thống kê thực tập sinh", text)
        self.assertIn("Nguyễn An", text)

    def test_large_dataset_keeps_detail_and_export_counts_aligned(self):
        users = [(number, 1, f"TTS {number}", f"tts{number}@example.test", "ThucTapSinh")
                 for number in range(100, 600)]
        profiles = [(number, number, 2, "Kỹ thuật", "DaDuyet", "2026-10-04 09:00:00")
                    for number in range(100, 600)]
        self.db.executemany("INSERT INTO NGUOI_DUNG VALUES (?, ?, ?, ?, ?)", users)
        self.db.executemany("INSERT INTO HO_SO_THUC_TAP VALUES (?, ?, ?, ?, ?, ?)", profiles)
        report = self.service.overview(InternStatisticsFilters(), 1, 100)
        self.assertEqual(report["total_items"], 503)
        self.assertEqual(len(list(self.service.iter_details(InternStatisticsFilters(), 503))), 503)
        response = reporting_routes.export_intern_statistics(
            request_for("HR"), "xlsx", program_id=None, department_id=None,
            date_from=None, date_to=None, school_id=None, major=None, search=None, db=self.db,
        )
        workbook = load_workbook(BytesIO(response.body), read_only=True)
        self.assertEqual(sum(isinstance(row[0], int) for row in workbook.active.values if row), 503)
        workbook.close()
        pdf = reporting_routes.export_intern_statistics(
            request_for("HR"), "pdf", program_id=None, department_id=None,
            date_from=None, date_to=None, school_id=None, major=None, search=None, db=self.db,
        )
        reader = PdfReader(BytesIO(pdf.body))
        self.assertGreater(len(reader.pages), 1)
        self.assertIn("TTS 100", reader.pages[-1].extract_text())


if __name__ == "__main__":
    unittest.main()
